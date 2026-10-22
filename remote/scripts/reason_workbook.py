#!/usr/bin/env python
"""Read Reason scope and remotable data from reason-remote_full.xlsx."""

import argparse
import json
import os
import sys

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
WORKBOOK = os.path.join(REPO, 'remote', 'mapping', 'reason-remote_full.xlsx')


def text(value):
    """Convert an Excel cell to the string form used by the map validator."""
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalized(value):
    return ''.join(char.lower() for char in value if char.isalnum())


def load_scope_index(path=WORKBOOK):
    try:
        import openpyxl
    except ImportError:
        raise SystemExit('openpyxl is required to read the Reason workbook: '
                         'install it with "python -m pip install openpyxl"')

    if not os.path.isfile(path):
        raise SystemExit('Reason workbook not found: %s' % path)

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if 'Index' not in workbook.sheetnames:
            raise ValueError('Reason workbook has no "Index" sheet')

        index_sheet = workbook['Index']
        rows = index_sheet.iter_rows(values_only=True)
        header = [text(value) for value in next(rows, ())]
        required = ('Device', 'Manufacturer', 'Model', 'Parameter Count')
        if any(name not in header for name in required):
            raise ValueError('Reason workbook Index sheet must contain: %s'
                             % ', '.join(required))
        columns = {name: header.index(name) for name in required}
        detail_sheets = [sheet for sheet in workbook.worksheets
                         if sheet.title != 'Index']
        index_entries = [row for row in rows
                         if len(row) > columns['Device']
                         and text(row[columns['Device']])]
        if len(index_entries) != len(detail_sheets):
            raise ValueError('Reason workbook Index has %d devices but %d '
                             'detail sheets'
                             % (len(index_entries), len(detail_sheets)))

        by_normalized_name = {}
        for sheet in detail_sheets:
            by_normalized_name.setdefault(normalized(sheet.title), []).append(sheet)

        scope_index = {}
        matched_sheets = set()
        for row in index_entries:
            device = text(row[columns['Device']])
            manufacturer = text(row[columns['Manufacturer']])
            model = text(row[columns['Model']])
            if not manufacturer or not model:
                raise ValueError('Reason workbook Index has an incomplete scope '
                                 'for %r' % device)

            name = normalized(device)
            exact = by_normalized_name.get(name, [])
            if len(exact) == 1:
                sheet = exact[0]
            else:
                matches = [candidate
                           for candidate in detail_sheets
                           if (normalized(candidate.title).startswith(name)
                               or name.startswith(normalized(candidate.title)))]
                if len(matches) != 1:
                    raise ValueError('Cannot uniquely match workbook device '
                                     '%r to a detail sheet' % device)
                sheet = matches[0]
            if sheet.title in matched_sheets:
                raise ValueError('Multiple Index devices match detail sheet %r'
                                 % sheet.title)
            matched_sheets.add(sheet.title)

            detail_rows = sheet.iter_rows(values_only=True)
            detail_header = [text(value) for value in next(detail_rows, ())]
            detail_columns = ('Remotable Item', 'Min Value', 'Max Value',
                              'Input Type', 'Output Type')
            if any(name not in detail_header for name in detail_columns):
                raise ValueError('Sheet %r is missing remotable columns'
                                 % sheet.title)
            positions = {name: detail_header.index(name)
                         for name in detail_columns}

            remotables = {}
            for detail in detail_rows:
                target = (text(detail[positions['Remotable Item']])
                          if len(detail) > positions['Remotable Item'] else '')
                if not target:
                    continue
                if target in remotables:
                    raise ValueError('Duplicate remotable %r in sheet %r'
                                     % (target, sheet.title))
                remotables[target] = {
                    'min': (text(detail[positions['Min Value']])
                            if len(detail) > positions['Min Value'] else ''),
                    'max': (text(detail[positions['Max Value']])
                            if len(detail) > positions['Max Value'] else ''),
                    'input': (text(detail[positions['Input Type']])
                              if len(detail) > positions['Input Type'] else ''),
                    'output': (text(detail[positions['Output Type']])
                               if len(detail) > positions['Output Type'] else ''),
                }

            count = row[columns['Parameter Count']]
            if count is not None and text(count).isdigit():
                if len(remotables) != int(count):
                    raise ValueError('Sheet %r has %d remotables; Index says %s'
                                     % (sheet.title, len(remotables), text(count)))

            key = '%s\t%s' % (manufacturer, model)
            if key in scope_index:
                raise ValueError('Duplicate scope %r in workbook Index' % key)
            scope_index[key] = {
                'manufacturer': manufacturer,
                'model': model,
                'display': None,
                'remotables': remotables,
                'factoryTargets': [],
                'sources': ['reason-remote_full.xlsx'],
            }

        return scope_index
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--targets-json', action='store_true',
                        help='write scopes and remotable names as JSON')
    args = parser.parse_args()

    scope_index = load_scope_index()
    if args.targets_json:
        result = {
            'scopes': [
                {
                    'key': key,
                    'manufacturer': entry['manufacturer'],
                    'model': entry['model'],
                    'targets': sorted(entry['remotables']),
                }
                for key, entry in sorted(scope_index.items())
            ],
        }
        json.dump(result, sys.stdout, ensure_ascii=False)
        sys.stdout.write('\n')
    else:
        print('Loaded %d scopes from %s'
              % (len(scope_index), os.path.relpath(WORKBOOK, REPO)))


if __name__ == '__main__':
    main()
